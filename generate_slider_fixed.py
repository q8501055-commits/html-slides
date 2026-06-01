import docx
import os
import re

doc_path = r"C:\Users\User\.openclaw\workspace\Humanoid_Robot_Tactile_Sensors_2026_Final.docx"
if not os.path.exists(doc_path):
    doc_path = r"C:\Users\User\.openclaw\workspace\Humanoid_Robot_Tactile_Sensors_2026_Ultimate.docx"
    
doc = docx.Document(doc_path)
paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

slides_data = []
current_slide = {"title": "簡報首頁", "content": []}

for p in paragraphs:
    if len(p) < 80 and (re.match(r'^[一二三四五六七八九十]、', p) or p.startswith("#") or "：" in p[:15] and len(p) < 30):
        if current_slide["title"] != "簡報首頁" or current_slide["content"]:
            slides_data.append(current_slide)
        title = p.replace("#", "").strip()
        current_slide = {"title": title, "content": []}
    else:
        current_slide["content"].append(p)

if current_slide["title"] or current_slide["content"]:
    slides_data.append(current_slide)

final_slides = []
for slide in slides_data:
    title = slide["title"]
    content = slide["content"]
    
    # Split content if it has too many paragraphs to prevent overflow
    # For text heavy, max 3 paragraphs per slide
    chunk_size = 3
    if not content:
        final_slides.append({"title": title, "content": []})
        continue
        
    for i in range(0, len(content), chunk_size):
        chunk = content[i:i+chunk_size]
        final_slides.append({
            "title": title if i == 0 else f"{title} (續)",
            "content": chunk
        })

html_template = """<!DOCTYPE html>
<html lang="zh-TW">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>2026 人形機器人觸覺感測器產業白皮書</title>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Orbitron:wght@400;700&family=Noto+Sans+TC:wght@300;400;700&display=swap');
        
        body, html {
            margin: 0; padding: 0; height: 100vh; width: 100vw;
            background: #050510;
            color: #0ff;
            font-family: 'Noto Sans TC', sans-serif;
            overflow: hidden;
        }

        .bg-grid {
            position: absolute; top: 0; left: 0; right: 0; bottom: 0;
            background-image: 
                linear-gradient(rgba(0, 255, 255, 0.05) 1px, transparent 1px),
                linear-gradient(90deg, rgba(0, 255, 255, 0.05) 1px, transparent 1px);
            background-size: 40px 40px;
            z-index: -1;
            transform: perspective(500px) rotateX(60deg) scale(2.5) translateY(-100px);
            animation: moveGrid 10s linear infinite;
        }

        @keyframes moveGrid {
            0% { background-position: 0 0; }
            100% { background-position: 0 40px; }
        }

        .slides-container {
            width: 100vw; height: 100vh;
            display: flex;
            transition: transform 0.6s cubic-bezier(0.25, 1, 0.5, 1);
        }

        .slide {
            min-width: 100vw; min-height: 100vh;
            display: flex; flex-direction: column;
            justify-content: center; align-items: center;
            position: relative;
            box-sizing: border-box;
            padding: 40px;
        }

        .glass-panel {
            background: rgba(10, 20, 40, 0.6);
            border: 1px solid rgba(0, 255, 255, 0.2);
            border-radius: 20px;
            padding: 50px;
            width: 90%;
            max-width: 1200px;
            max-height: 85vh;
            overflow-y: auto;
            box-shadow: 0 0 30px rgba(0, 255, 255, 0.1), inset 0 0 20px rgba(0, 255, 255, 0.05);
            backdrop-filter: blur(15px);
            animation: float 6s ease-in-out infinite;
            text-align: left;
        }
        
        .glass-panel::-webkit-scrollbar { width: 8px; }
        .glass-panel::-webkit-scrollbar-track { background: rgba(0, 255, 255, 0.05); }
        .glass-panel::-webkit-scrollbar-thumb { background: rgba(0, 255, 255, 0.3); border-radius: 4px; }

        @keyframes float {
            0% { transform: translateY(0px) rotateX(0deg); }
            50% { transform: translateY(-10px) rotateX(1deg); }
            100% { transform: translateY(0px) rotateX(0deg); }
        }

        h1, h2, h3 {
            font-family: 'Orbitron', 'Noto Sans TC', sans-serif;
            color: #fff;
            text-shadow: 0 0 10px #0ff, 0 0 20px #00aaff;
            letter-spacing: 1px;
            margin-top: 0;
            border-bottom: 1px solid rgba(0,255,255,0.3);
            padding-bottom: 15px;
        }
        
        h1 { font-size: 2.5em; text-align: center; border-bottom: none; }
        h2 { font-size: 2em; color: #0ff; }

        p, li {
            font-size: 1.1em;
            line-height: 1.8;
            color: #e0ffff;
            text-shadow: 0 0 5px rgba(0,255,255,0.3);
            margin-bottom: 15px;
        }

        .controls {
            position: fixed; bottom: 20px; left: 50%;
            transform: translateX(-50%);
            z-index: 1000;
            display: flex;
            gap: 20px;
        }

        button {
            background: rgba(0, 255, 255, 0.1);
            border: 1px solid #0ff;
            color: #0ff;
            padding: 10px 25px;
            font-family: 'Orbitron', sans-serif;
            font-size: 1.2em;
            cursor: pointer;
            border-radius: 5px;
            transition: all 0.3s ease;
            box-shadow: 0 0 10px rgba(0, 255, 255, 0.2);
        }

        button:hover {
            background: rgba(0, 255, 255, 0.3);
            box-shadow: 0 0 20px rgba(0, 255, 255, 0.8);
            text-shadow: 0 0 5px #fff;
        }
        
        .progress-bar {
            position: fixed; top: 0; left: 0; height: 4px;
            background: #0ff;
            box-shadow: 0 0 10px #0ff;
            z-index: 1000;
            transition: width 0.4s ease;
        }

        .img-placeholder {
            margin: 20px auto;
            display: block;
            max-height: 35vh;
            border-radius: 8px;
            box-shadow: 0 0 15px rgba(0,255,255,0.4);
            border: 1px solid rgba(0,255,255,0.2);
        }
        
        /* Highlight tags */
        .highlight {
            color: #f0f;
            text-shadow: 0 0 10px #f0f;
            font-weight: bold;
        }
    </style>
</head>
<body>
    <div class="bg-grid"></div>
    <div class="progress-bar" id="progress"></div>
    
    <div class="slides-container" id="slides-container">
"""

html_template += """
        <div class="slide">
            <div class="glass-panel" style="text-align: center;">
                <h1 style="font-size: 3.5em; margin-bottom: 20px;">人形機器人觸覺感測器</h1>
                <h2 style="border: none;">2026 產業深度白皮書與技術佈局</h2>
                <p style="text-align: center; margin-top: 30px;">完整技術細節 / 供應鏈分析 / 未來戰略</p>
            </div>
        </div>
"""

img_index = 0
for i, slide in enumerate(final_slides):
    html_template += f"""
        <div class="slide">
            <div class="glass-panel">
                <h2>{slide['title']}</h2>
"""
    for p in slide['content']:
        if "：" in p:
            parts = p.split("：", 1)
            p_html = f"<span class='highlight'>{parts[0]}：</span>{parts[1]}"
        else:
            p_html = p
        html_template += f"<p>{p_html}</p>\n"
    
    if i > 0 and i % 4 == 0:
        img_src = f"assets/sensor_{img_index % 3}.jpg"
        html_template += f'<div style="text-align:center;"><img src="{img_src}" class="img-placeholder" alt="Sensor Tech" onerror="this.style.display=\'none\'"></div>'
        img_index += 1

    html_template += """
            </div>
        </div>
"""

html_template += """
    </div>
    <div class="controls">
        <button onclick="prevSlide()">&#9664; PREV</button>
        <button onclick="nextSlide()">NEXT &#9654;</button>
    </div>

    <script>
        let currentSlide = 0;
        const slides = document.querySelectorAll('.slide');
        const totalSlides = slides.length;
        const container = document.getElementById('slides-container');
        const progress = document.getElementById('progress');

        function updateSlide() {
            container.style.transform = `translateX(-${currentSlide * 100}vw)`;
            progress.style.width = `${((currentSlide + 1) / totalSlides) * 100}%`;
        }

        function nextSlide() {
            if (currentSlide < totalSlides - 1) {
                currentSlide++;
                updateSlide();
            }
        }

        function prevSlide() {
            if (currentSlide > 0) {
                currentSlide--;
                updateSlide();
            }
        }

        document.addEventListener('keydown', (e) => {
            if (e.key === 'ArrowRight' || e.key === ' ' || e.key === 'PageDown') nextSlide();
            if (e.key === 'ArrowLeft' || e.key === 'PageUp') prevSlide();
        });
        
        updateSlide();
    </script>
</body>
</html>
"""

with open(r"C:\Users\User\.openclaw\workspace\html-slides-repo\Humanoid_Robot_Tactile_Sensors_2026_Final.html", "w", encoding="utf-8") as f:
    f.write(html_template)

print(f"Generated successfully. Slides count: {len(final_slides) + 1}")
